#!/usr/bin/env python3
"""Fill missing US mass transit agency HQ address/phone fields."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent
SRC = Path("/Users/jeffcoleman/Downloads/NexCort_iQ_US_Mass_Transit_Agency_Outreach.xlsx")
OUT = Path(
    "/Users/jeffcoleman/Downloads/NexCort_iQ_US_Mass_Transit_Agency_Outreach_FILLED.xlsx"
)


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not str(v).strip())


def norm_phone(p) -> str | None:
    if not p:
        return None
    s = str(p).strip()
    digits = re.sub(r"\D", "", s)
    if len(digits) == 11 and digits[0] == "1":
        return f"1-{digits[1:4]}-{digits[4:7]}-{digits[7:]}"
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return s


def main() -> None:
    rows: list[dict] = []
    for name in [
        "transit_batch_1_30.json",
        "transit_batch_31_60.json",
        "transit_batch_61_90.json",
        "transit_batch_91_117.json",
    ]:
        arr = json.loads((ROOT / name).read_text())
        print(name, len(arr))
        rows.extend(arr)
    by_id = {int(x["id"]): x for x in rows}
    (ROOT / "transit_all.json").write_text(
        json.dumps({str(k): v for k, v in by_id.items()}, indent=2)
    )
    print("research ids", len(by_id))
    print(
        "missing research for",
        sorted(set(range(1, 118)) - set(by_id) - {6, 9, 21, 34}),
    )

    wb = load_workbook(SRC)
    ws = wb["Transit Agency Outreach"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    stats = {"addr": 0, "phone": 0, "src": 0, "rows": 0}

    for r in range(2, ws.max_row + 1):
        rid = ws.cell(r, h["#"]).value
        data = by_id.get(int(rid)) if rid is not None else None
        if not data:
            continue
        changed = False

        def set_blank(col: str, val, counter: str | None = None) -> None:
            nonlocal changed
            if blank(val):
                return
            cell = ws.cell(r, h[col])
            if blank(cell.value):
                cell.value = val
                changed = True
                if counter:
                    stats[counter] += 1

        set_blank(
            "Headquarters / administrative address", data.get("address"), "addr"
        )
        set_blank("Phone to call", norm_phone(data.get("phone")), "phone")
        set_blank("Contact source URL", data.get("source_url"), "src")

        src_cell = ws.cell(r, h["Contact source URL"])
        if src_cell.value and "wwww.septa" in str(src_cell.value):
            src_cell.value = data.get("source_url") or "https://www.septa.org/about"
            changed = True

        phone_cell = ws.cell(r, h["Phone to call"])
        if data.get("notes") and not blank(phone_cell.value):
            if phone_cell.comment is None:
                phone_cell.comment = Comment(
                    f"Confidence: {data.get('confidence')}. "
                    f"{(data.get('notes') or '')[:400]}",
                    "NexCort research",
                )

        if changed:
            stats["rows"] += 1

    phones = sum(
        1
        for r in range(2, ws.max_row + 1)
        if not blank(ws.cell(r, h["Phone to call"]).value)
    )
    rm = wb["Read Me & Summary"]
    for r in range(1, rm.max_row + 1):
        key = rm.cell(r, 1).value
        if key == "Verified published phone numbers":
            rm.cell(r, 2).value = phones
        if key == "Contact limitations":
            rm.cell(r, 2).value = (
                "HQ addresses and main switchboards filled from official agency contact "
                "pages (2026-10-11). Named decision makers still require verification. "
                "Never use emergency/transit-police numbers for sales."
            )
        if key == "Data currency":
            rm.cell(r, 2).value = (
                "Compiled October 2026; HQ phones/addresses filled 2026-10-11 — "
                "independently confirm each contact before outreach."
            )
    rm.cell(rm.max_row + 1, 1).value = "Fill method"
    rm.cell(rm.max_row, 2).value = (
        "Official agency Contact Us / About HQ pages; admin or customer switchboard "
        "preferred; source URLs in Contact source URL; notes in phone cell comments."
    )

    wb.save(OUT)
    wb.save(SRC)

    wb3 = load_workbook(SRC)
    ws3 = wb3["Transit Agency Outreach"]
    headers = [ws3.cell(1, c).value for c in range(1, ws3.max_column + 1)]
    blanks = {n: 0 for n in headers}
    still = []
    for r in range(2, ws3.max_row + 1):
        for c, n in enumerate(headers, 1):
            if blank(ws3.cell(r, c).value):
                blanks[n] += 1
        if blank(ws3.cell(r, headers.index("Phone to call") + 1).value) or blank(
            ws3.cell(
                r, headers.index("Headquarters / administrative address") + 1
            ).value
        ):
            still.append(
                (
                    ws3.cell(r, 1).value,
                    ws3.cell(r, 4).value,
                )
            )
    print("STATS", stats)
    print("REMAINING", blanks)
    print("STILL INCOMPLETE", still, "count", len(still))
    for rid in (1, 5, 12, 35, 91, 117):
        for r in range(2, ws3.max_row + 1):
            if ws3.cell(r, 1).value == rid:
                print(rid, [ws3.cell(r, c).value for c in range(1, 13)])
                break
    print("OUT", OUT)


if __name__ == "__main__":
    main()
