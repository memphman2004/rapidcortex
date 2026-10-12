#!/usr/bin/env python3
"""Fill missing 50-state community college system outreach fields."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent
SRC = Path("/Users/jeffcoleman/Downloads/NexCort_iQ_50_State_Community_College_Systems.xlsx")
OUT = Path("/Users/jeffcoleman/Downloads/NexCort_iQ_50_State_Community_College_Systems_FILLED.xlsx")
ENROLL_COL = "Enrollment "


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and not str(v).strip())


def norm_phone(p) -> str | None:
    if not p:
        return None
    s = str(p).strip()
    digits = re.sub(r"\D", "", s)
    if len(digits) == 11 and digits[0] == "1":
        digits = digits[1:]
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return s


def zip5(z) -> str | None:
    if not z:
        return None
    return str(z).split("-")[0][:5]


def main() -> None:
    contacts: list[dict] = []
    for name in [
        "cc_batch_al_ky.json",
        "cc_batch_la_nm.json",
        "cc_batch_nd_wy.json",
    ]:
        contacts.extend(json.loads((ROOT / name).read_text()))
    by_state = {c["state"]: c for c in contacts}
    enroll11 = {
        r["state"]: r for r in json.loads((ROOT / "cc_enrollment_11.json").read_text())
    }

    wb = load_workbook(SRC)
    ws = wb["50 State Systems"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    stats = {"contact_rows": 0, "enroll": 0, "rows": 0}

    for r in range(2, ws.max_row + 1):
        state = ws.cell(r, h["State"]).value
        data = by_state.get(state)
        e11 = enroll11.get(state)
        changed = False

        def set_blank(col: str, val) -> bool:
            nonlocal changed
            if blank(val):
                return False
            cell = ws.cell(r, h[col])
            if blank(cell.value):
                cell.value = val
                changed = True
                return True
            return False

        if data:
            office = ws.cell(r, h["Named Point of Contact / Office"])
            if blank(office.value) or str(office.value).strip() == "Research needed":
                if data.get("named_office"):
                    office.value = data["named_office"]
                    changed = True
                    stats["contact_rows"] += 1
            set_blank("Contact Title", data.get("contact_title"))
            set_blank("Street / Mailing Address", data.get("street"))
            set_blank("City", data.get("city"))
            set_blank("ST", data.get("st"))
            set_blank("ZIP", zip5(data.get("zip")))
            set_blank("Phone to Call", norm_phone(data.get("phone")))

            enc = ws.cell(r, h[ENROLL_COL])
            if blank(enc.value) and data.get("enrollment") is not None:
                enc.value = int(data["enrollment"])
                changed = True
                stats["enroll"] += 1
                yr = data.get("enrollment_year")
                src = data.get("enrollment_source")
                notes = (data.get("notes") or "")[:300]
                enc.comment = Comment(
                    f"Enrollment year/context: {yr}. Source: {src} | {notes}",
                    "NexCort research",
                )

            phone_cell = ws.cell(r, h["Phone to Call"])
            if data.get("source_url") and not blank(phone_cell.value):
                if phone_cell.comment is None:
                    phone_cell.comment = Comment(
                        f"Contact source: {data['source_url']}",
                        "NexCort research",
                    )

        if e11 and e11.get("enrollment") is not None:
            enc = ws.cell(r, h[ENROLL_COL])
            # Prefer dedicated enroll11 figures for the 11 pre-contacted states
            if blank(enc.value) or state in enroll11:
                prev = enc.value
                enc.value = int(e11["enrollment"])
                if blank(prev):
                    stats["enroll"] += 1
                if prev != enc.value:
                    changed = True
                yr = e11.get("enrollment_year")
                src = e11.get("enrollment_source")
                notes = (e11.get("notes") or "")[:350]
                enc.comment = Comment(
                    f"Enrollment year/context: {yr}. Source: {src} | {notes}",
                    "NexCort research",
                )
            if state == "Florida":
                set_blank(
                    "Street / Mailing Address",
                    e11.get("florida_street")
                    or (data.get("street") if data else None),
                )
                set_blank(
                    "ZIP",
                    zip5(
                        e11.get("florida_zip")
                        or (data.get("zip") if data else None)
                    ),
                )

        if changed:
            stats["rows"] += 1

    rm = wb["Read Me"]
    for r in range(1, rm.max_row + 1):
        key = rm.cell(r, 1).value
        if key == "Contact status":
            rm.cell(r, 2).value = (
                "System/association HQ phones populated from official sites "
                "(filled 2026-10-11). Confirm before outreach."
            )
        if key == "Enrollments & Crime":
            rm.cell(r, 2).value = (
                "Statewide/system headcounts filled from IR/board/association reports "
                "(see cell comments — Fall census vs annual unduplicated definitions vary). "
                "Clery join still pending."
            )
        if key == "Data date":
            rm.cell(r, 2).value = (
                "2026-10-11 (contacts + enrollment filled); confirm phone numbers "
                "and staff before outreach."
            )
    rm.cell(rm.max_row + 1, 1).value = "Fill method"
    rm.cell(rm.max_row, 2).value = (
        "Official system/association HQ pages for address/phone; enrollment from "
        "IR/board reports (definitions vary by state — see comments)."
    )

    wb.save(OUT)
    wb.save(SRC)

    wb3 = load_workbook(SRC)
    ws3 = wb3["50 State Systems"]
    headers = [ws3.cell(1, c).value for c in range(1, ws3.max_column + 1)]
    blanks = {name: 0 for name in headers}
    for r in range(2, ws3.max_row + 1):
        for c, name in enumerate(headers, 1):
            if blank(ws3.cell(r, c).value):
                blanks[name] += 1
    print("STATS", stats)
    print("REMAINING", blanks)
    for st in [
        "Alabama",
        "Florida",
        "California",
        "Virginia",
        "Wyoming",
        "Montana",
        "Texas",
    ]:
        for r in range(2, ws3.max_row + 1):
            if ws3.cell(r, 1).value == st:
                print(st, [ws3.cell(r, c).value for c in range(1, 15)])
                break
    print("OUT", OUT)


if __name__ == "__main__":
    main()
