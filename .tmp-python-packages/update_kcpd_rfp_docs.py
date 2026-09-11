#!/usr/bin/env python3
"""Update KCPD Rapid Cortex reply documents: accurate facts + 35% price cut."""

from __future__ import annotations

import shutil
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

COMPLETE = Path(
    "/Volumes/Mac Mini/Business Documents/RFP:RFI/5 Kansas City Police 30 Sep/Complete"
)
XLSX = COMPLETE / "RC_KCPD_Pricing_Schedule.xlsx"
DOCX = COMPLETE / "RC_KCPD_Proposal_Scope_of_Services.docx"


def disc(n: int | float | Decimal) -> int:
    """35% reduction, round half up to whole dollars."""
    return int((Decimal(str(n)) * Decimal("0.65")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


# --- priced lines (original → discounted) ---
P = {
    1: disc(648000),  # 421,200
    3: disc(90000),  # 58,500
    5: disc(132000),  # 85,800
    6: disc(66000),  # 42,900
    7: disc(90000),  # 58,500
    8: disc(30000),  # 19,500
    11: disc(278750),  # 181,188
    14: disc(65000),  # 42,250
    16: disc(32500),  # 21,125
    18: disc(42000),  # 27,300
    19: disc(210000),  # 136,500
    23: disc(101250),  # 65,813
    24: disc(477500),  # 310,375
    25: disc(240000),  # 156,000
}

ANNUAL = P[1] + P[3] + P[5] + P[6] + P[7] + P[8] + P[18] + P[19] + P[25]
ONE_TIME = P[11] + P[14] + P[16] + P[23] + P[24]
YEAR1 = ANNUAL + ONE_TIME
CAD_MAINT = disc(27500)  # 17,875
RMS_MAINT = disc(11250)  # 7,313
YEAR2 = ANNUAL + CAD_MAINT + RMS_MAINT
CMD_BASE = disc(480000)  # 312,000
CALL_ASSIST = disc(168000)  # 109,200

ENTITY = "Apps on Demand LLC d/b/a Rapid Cortex"
CONTACT = "Dr. Jeffrey W. Coleman Jr., Founder & CEO"
HQ = "Columbus, Georgia 31901"
EMAIL = "press@rapidcortex.us"
WEB = "rapidcortex.us"

CAD_ADAPTERS = (
    "Motorola PremierOne, Tyler New World, Hexagon I/CAD, CentralSquare, "
    "Zetron, Mark43, and Versaterm"
)

CAD_NOTE = (
    f"Includes (35% proposal discount applied): Advanced Discovery ($9,750) + "
    f"Field Mapping ($9,750) + Vendor Coordination ($11,375) + Sandbox Testing ($16,250) + "
    f"Assisted Write-Back ($81,250) + Audit Logging ($13,000) + Rollback Planning ($6,500) + "
    f"Integration Mapping ($9,750) + Program Management ($23,563). "
    f"Year 2+ maintenance: ${CAD_MAINT:,}/yr billed separately. "
    f"Write-back is human-reviewed (assisted) and fail-closed until KCPD UAT."
)

RMS_NOTE = (
    f"Motorola Records read integration + draft report write-back (Tier 2). "
    f"Pre-populated incident draft reports for eligible call types; human review required "
    f"before submission. Live Records filing stays fail-closed until validated. "
    f"Year 2+ maintenance: ${RMS_MAINT:,}/yr billed separately."
)

LINE1_NOTE = (
    f"Command platform base (${CMD_BASE:,}/yr) + Call Assist voice AI module "
    f"(${CALL_ASSIST:,}/yr). Includes telephony engine, dynamic questioning, distress "
    f"detection, emergency escalation, confidence scoring, and AI disclosure. "
    f"Non-emergency line only — emergencies transfer immediately to a live call taker. "
    f"35% proposal discount applied."
)

FOOTER_NOTES = (
    f"Notes: All pricing in U.S. dollars. A 35% strategic proposal discount is applied to "
    f"every priced line item below Rapid Cortex list rates. Annual subscription rates locked "
    f"for 36 months from contract execution. Year 4+ pricing subject to maximum 4% annual "
    f"CPI adjustment. Travel expenses for onsite services billed at cost per travel policy "
    f"in the MSA. CAD integration maintenance (${CAD_MAINT:,}/yr) and RMS maintenance "
    f"(${RMS_MAINT:,}/yr) begin in Year 2. Official bid line pricing is also submitted via "
    f"the KCPD online bid platform Line Item tab. This pricing schedule is submitted as an "
    f"optional supplemental attachment per RFP Section 13. Contracting entity: {ENTITY}, "
    f"{HQ}."
)


def money(n: int) -> str:
    return f"${n:,}"


def backup(path: Path) -> None:
    bak = path.with_suffix(path.suffix + ".pre-35pct.bak")
    if not bak.exists():
        shutil.copy2(path, bak)


def copy_cell_style(src: Cell, dst: Cell) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def update_xlsx() -> None:
    backup(XLSX)
    wb = load_workbook(XLSX)
    ps = wb["Pricing Schedule"]

    ps["B3"] = ENTITY
    ps["D7"] = P[1]
    ps["F7"] = LINE1_NOTE
    ps["D9"] = P[3]
    ps["F10"] = (
        "Included in Bid Line 3. Auto-classifies calls across all 12 RFP categories "
        "plus configurable additional types. Admin-adjustable rules without vendor involvement."
    )
    ps["D11"] = P[5]
    ps["F11"] = (
        "NLP/LLM conversation management, advanced confidence scoring, AI incident summarization, "
        "model redundancy (primary/secondary/tertiary), intent detection, Amazon Comprehend "
        "sentiment plus Contact Lens / lexical distress detection, and dynamic questioning. "
        "Not a custom acoustic emotion CNN."
    )
    ps["D12"] = P[6]
    ps["F12"] = (
        "English + Spanish voice/transcription; live interpreter-bridge routing (tenant config); "
        "TTY/TDD via Amazon Connect attributes with SMS fallback; language auto-detection; "
        "real-time translation (included in Command plan; this line covers enhanced Tier 3 "
        "transcription and TTY/interpreter extras)."
    )
    ps["D13"] = P[7]
    ps["F13"] = (
        "Voice included in Line 1. SMS link generation. Caller video upload (async) attaches to "
        "existing Rapid Cortex media. Live caller video streaming (WebRTC) available as a "
        "configured channel. All media logged with audit trail."
    )
    ps["D14"] = P[8]
    ps["F14"] = (
        "Callback number capture, after-hours callback offer, and caller status notifications "
        "via SMS. Automated outbound callback campaigns are configured per KCPD call type and "
        "priority during implementation."
    )
    ps["D17"] = P[11]
    ps["F17"] = CAD_NOTE
    ps["F18"] = (
        "Removed per RFP 2026-0010 Addendum. RapidSOS Radius/Clearinghouse integration is "
        "available to KCPD at no additional integration fee (location/supplemental data at call time)."
    )
    ps["F19"] = (
        "Amazon Connect SIP/VoIP bridge to KCPD's non-emergency line. ANI/ALI passthrough "
        "preserved. Queue management and overflow routing per KCPD rules. Does not replace "
        "KCPD's 911 CPE/phone system. No telephony hardware replacement required."
    )
    ps["D20"] = P[14]
    ps["F20"] = RMS_NOTE
    ps["F21"] = (
        "Platform mapping: caller-stated location, address capture, and optional RapidSOS "
        "location candidate. Zone/beat assignment from KCPD-configured GIS/address data. "
        "No additional GIS licensing required."
    )
    ps["D22"] = P[16]
    ps["F22"] = (
        "CJIS Security Policy–aligned controls review and documentation package. MFA, TLS 1.2+, "
        "AES-256 encryption, RBAC, and audit logging are included in the base platform. Rapid Cortex "
        "does not issue FBI CJIS, SOC 2, or FedRAMP certification — KCPD maps controls to its CJIS program."
    )
    ps["F23"] = (
        "Core cybersecurity posture included (AWS WAF, GuardDuty, CloudTrail, Security Hub, "
        "tenant isolation, Secrets Manager). Coordinate with KCPD IT (816) 234-5286. "
        "Annual third-party pen test results available under NDA. Not a staffed customer SOC."
    )
    ps["D24"] = P[18]
    ps["D25"] = P[19]
    ps["F25"] = (
        "Supervisor QA tools, scorecards, coaching notes, and team dashboards included in Command. "
        "This line covers AI QA scoring (KCPD rubric), call playback, keyword search, false-transfer "
        "tracking, escalation statistics, and performance dashboards."
    )
    ps["F28"] = (
        "Platform targets: 99.99% availability with SLA credits for verified downtime per the MSA; "
        "<2s AI response (95th percentile); <500ms API latency (99th percentile); 95% transcription "
        "accuracy target; 95% location recognition; 24/7 monitoring; RTO < 4 hrs; RPO ≤ 15 min. "
        "Monthly SLA reporting."
    )
    ps["D29"] = P[23]
    ps["D30"] = P[24]
    ps["D31"] = P[25]
    ps["F32"] = (
        f"Live scenario-based demonstration at no additional cost (abandoned vehicle, noise, theft, "
        f"welfare check, parking, Spanish, CarFax, mid-call emergency). Contracting entity: {ENTITY}, "
        f"{HQ}. CAD portability via CAD abstraction layer ({CAD_ADAPTERS}). Three law enforcement "
        f"references provided in the Reference Information Sheet."
    )

    ps["D34"] = ANNUAL
    ps["D35"] = ONE_TIME
    ps["D36"] = YEAR1
    ps["D37"] = YEAR2
    ps["A40"] = FOOTER_NOTES

    ref = wb["Bid Line Entry Reference"]
    ref["D4"] = P[1]
    ref["E4"] = "Annual — Command Large T3 + Call Assist (35% proposal discount)"
    ref["D6"] = P[3]
    ref["D8"] = P[5]
    ref["D9"] = P[6]
    ref["D10"] = P[7]
    ref["D11"] = P[8]
    ref["D14"] = P[11]
    ref["E14"] = "One-time — PremierOne assisted write-back (human-reviewed; UAT gated)"
    ref["D17"] = P[14]
    ref["E17"] = "One-time — Motorola Records read + draft write (fail-closed until UAT)"
    ref["D19"] = P[16]
    ref["D21"] = P[18]
    ref["D22"] = P[19]
    ref["D26"] = P[23]
    ref["D27"] = P[24]
    ref["D28"] = P[25]
    ref["A1"] = (
        "RAPID CORTEX — BID LINE ENTRY REFERENCE  |  KCPD Line Item Tab  |  "
        "35% proposal discount applied  |  " + ENTITY
    )

    wb.save(XLSX)
    print("Excel updated:")
    print(f"  Annual {ANNUAL:,}  One-time {ONE_TIME:,}  Year 1 {YEAR1:,}  Year 2+ {YEAR2:,}")


def set_cell_text(cell, text: str) -> None:
    """Replace a table cell's text while keeping the first run's formatting."""
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.text = text
        return
    first = paragraphs[0]
    if first.runs:
        first.text = text
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.add_run(text)
    for extra in paragraphs[1:]:
        extra.clear()


def replace_in_paragraph(p, old: str, new: str) -> bool:
    if old not in p.text:
        return False
    if p.runs:
        # Prefer a single-run hit
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
        # Text split across runs — rewrite first run, clear the rest
        full = p.text.replace(old, new)
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ""
        return True
    p.add_run(new if not p.text else p.text.replace(old, new))
    return True


def replace_all(doc: Document, pairs: list[tuple[str, str]]) -> int:
    n = 0
    for p in doc.paragraphs:
        for old, new in pairs:
            if old in p.text:
                if replace_in_paragraph(p, old, new):
                    n += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old, new in pairs:
                        if old in p.text:
                            if replace_in_paragraph(p, old, new):
                                n += 1
    for section in doc.sections:
        for hdr in (section.header, section.footer):
            for p in hdr.paragraphs:
                for old, new in pairs:
                    if old in p.text:
                        if replace_in_paragraph(p, old, new):
                            n += 1
    return n


def find_table(doc: Document, header_contains: str):
    for t in doc.tables:
        if not t.rows:
            continue
        head = " | ".join(c.text for c in t.rows[0].cells)
        if header_contains.lower() in head.lower():
            return t
    return None


def update_docx() -> None:
    backup(DOCX)
    doc = Document(str(DOCX))

    # Cover / identity
    pairs: list[tuple[str, str]] = [
        ("Sharen.Marquez@kcpd.org", "Sharene.Marquez@kcpd.org"),
        (
            "Rapid Cortex is a public safety intelligence and decision-support technology company headquartered in the United States. We design, develop, and operate real-time AI-assisted platforms for 911 dispatch centers, emergency communications centers, law enforcement agencies, fire rescue, EMS, and emergency management organizations.",
            f"{ENTITY} designs, develops, and operates Rapid Cortex — a real-time AI-assisted public safety intelligence platform for 911 dispatch centers, emergency communications centers, law enforcement, fire, EMS, and emergency management. The company is headquartered in {HQ}. Primary contact: {CONTACT} ({EMAIL}).",
        ),
        (
            "Our platform currently serves operational PSAP customers across multiple U.S. jurisdictions. All customer deployments remain active production systems — Rapid Cortex does not operate pilot-only programs counted as references.",
            "Rapid Cortex is purpose-built for emergency communications. Call Assist, the module proposed here, answers and triages non-emergency calls, transfers emergencies immediately to a live telecommunicator, and never replaces CAD, 911 CPE, or dispatcher authority. References are provided in the Reference Information Sheet.",
        ),
        (
            "Rapid Cortex operates exclusively in the public safety vertical. Our team includes former emergency communications directors, PSAP technology administrators, public safety software engineers, and compliance professionals with direct CJIS experience. We do not deploy consumer AI products and do not adapt general-purpose software for public safety use — our architecture is designed ground-up for the operational, legal, and compliance requirements of PSAPs and law enforcement agencies.",
            f"Rapid Cortex operates exclusively in the public safety vertical. Architecture, RBAC, audit logging, and tenant isolation are designed ground-up for PSAP and law enforcement requirements. We do not deploy consumer AI products. Cybersecurity claims in this proposal are CJIS-aligned technical controls — not FBI CJIS, SOC 2 Type II, or FedRAMP certification.",
        ),
        (
            "Currently certified connectors include: Motorola PremierOne, Tyler New World CAD, Hexagon I/CAD, Central Square CAD, and CommandAI. Should KCPD transition to a new CAD platform, Rapid Cortex migrates the integration without a full platform re-deployment, protecting KCPD's investment.",
            f"The CAD abstraction layer includes adapters for {CAD_ADAPTERS}. KCPD's Motorola PremierOne deployment is a scoped integration project (discovery, sandbox, UAT) — not a checkbox. Should KCPD transition CAD platforms, Rapid Cortex migrates the adapter without a full platform re-deployment. Production write-back remains fail-closed until KCPD UAT.",
        ),
        (
            "Rapid Cortex uses a CAD abstraction layer supporting Motorola PremierOne, Tyler New World, Hexagon CAD, Central Square, and additional platforms. KCPD is not locked in to any single CAD vendor.",
            f"Rapid Cortex uses a CAD abstraction layer with adapters for {CAD_ADAPTERS}. KCPD is not locked to a single CAD vendor. Live write-back requires write-back flags plus agency UAT.",
        ),
        (
            "SOC 2 Type II alignment, incident response procedures, security event monitoring, AWS security services integration (GuardDuty, Security Hub, CloudTrail), and contractual security SLAs.",
            "AWS security services (GuardDuty, Security Hub, CloudTrail, WAF), incident-response procedures, and security-event monitoring. Rapid Cortex documents CJIS-aligned controls; it does not hold SOC 2 Type II, CJIS, or FedRAMP certification.",
        ),
        (
            "Rapid Cortex operates its platform in alignment with the FBI CJIS Security Policy, Version 5.9. The following summary identifies key compliance elements relevant to KCPD's requirements:",
            "Rapid Cortex operates with CJIS Security Policy–aligned technical controls (FBI CJIS Security Policy). The following summary identifies key control areas relevant to KCPD. Rapid Cortex does not issue FBI approval; KCPD maps these controls to its CJIS program.",
        ),
        (
            "AWS GovCloud-compatible architecture; data residency within United States; no CJIS data leaves U.S. jurisdiction.",
            "AWS US-region architecture; data residency within the United States; no KCPD operational data leaves U.S. jurisdiction. GovCloud can be scoped by SOW if KCPD requires it.",
        ),
        (
            "The following law enforcement and PSAP agency references are provided in support of Bid Line 26 (Vendor Qualifications). KCPD evaluators are welcome to contact any reference directly. All agencies listed operate Rapid Cortex in active production.",
            "The following law enforcement and PSAP agency references are provided in support of Bid Line 26 (Vendor Qualifications). KCPD evaluators are welcome to contact any reference directly. Contact details are in the Reference Information Sheet submitted separately.",
        ),
        (
            "Questions: Contact your Rapid Cortex Account Executive",
            f"Questions: {CONTACT}  |  {EMAIL}  |  {WEB}",
        ),
        # Pricing — original strings
        ("$648,000", money(P[1])),
        ("$480K/yr", money(CMD_BASE) + "/yr"),
        ("$168K/yr", money(CALL_ASSIST) + "/yr"),
        ("$90,000", money(P[3])),  # also Line 7 — handled by later unique strings where needed
        ("$132,000", money(P[5])),
        ("$66,000", money(P[6])),
        ("$30,000", money(P[8])),
        ("$278,750", money(P[11])),
        ("$65,000", money(P[14])),
        ("$11,250/yr", money(RMS_MAINT) + "/yr"),
        ("$32,500", money(P[16])),
        ("$42,000", money(P[18])),
        ("$210,000", money(P[19])),
        ("$101,250", money(P[23])),
        ("$477,500", money(P[24])),
        ("$240,000", money(P[25])),
        ("$955,000", money(ONE_TIME)),
        ("$1,548,000", money(ANNUAL)),
        ("$2,503,000", money(YEAR1)),
        ("$1,586,750", money(YEAR2)),
        (
            "Advanced Discovery ($15K) + Vendor Coordination ($17.5K) + Assisted Write-Back Tier 2 ($125K) + Sandbox ($25K) + Audit Logging ($20K) + Rollback Planning ($10K) + Mapping ($15K) + 15% maint. yr2+ ($27.5K/yr noted separately)",
            f"Discovery, mapping, Motorola coordination, sandbox, assisted write-back (human-reviewed), audit logging, rollback, and program management. Year 2+ maintenance {money(CAD_MAINT)}/yr. 35% proposal discount applied.",
        ),
        (
            "The pricing schedule below maps directly to the 26 bid lines in RFP 2026-0010 Addendum 4. Annual subscription pricing reflects the full KCPD deployment on the Rapid Cortex Command Large T3 plan (51–65 telecommunicator seats). One-time fees are charged at contract execution. Year 2+ recurring reflects ongoing annual costs plus integration maintenance.",
            f"The pricing schedule below maps directly to the 26 bid lines in RFP 2026-0010 Addendum 4. Annual subscription pricing reflects the full KCPD deployment on Rapid Cortex Command Large T3 (51–65 telecommunicator seats) plus Call Assist. A 35% strategic proposal discount is applied to every priced line. One-time fees are charged at contract execution. Year 2+ recurring reflects ongoing annual costs plus integration maintenance ({money(CAD_MAINT)} CAD + {money(RMS_MAINT)} RMS).",
        ),
        (
            "All pricing is in U.S. dollars. Annual subscription rates are locked for 36 months from contract execution. Year 4+ pricing subject to a maximum annual CPI adjustment not to exceed 4%. Pricing does not include travel expenses for onsite services; travel is billed at cost per the travel policy included in the Master Service Agreement.",
            f"All pricing is in U.S. dollars and reflects a 35% strategic proposal discount from Rapid Cortex list rates. Annual subscription rates are locked for 36 months from contract execution. Year 4+ pricing subject to a maximum annual CPI adjustment not to exceed 4%. Pricing does not include travel expenses for onsite services; travel is billed at cost per the travel policy included in the Master Service Agreement. Contracting entity: {ENTITY}.",
        ),
        (
            "Dedicated KCPD production environment — no resource sharing with other agencies during peak demand.",
            "Dedicated KCPD tenant isolation (agencyId-scoped data, no cross-tenant access). Compute is multi-tenant AWS with per-agency data isolation.",
        ),
        (
            "Custom report templates can be configured by KCPD administrators.",
            "Custom report templates can be configured by KCPD administrators. Geographic heat maps and full analytics exports are included in Command; Call Assist v1 surfaces session counts, transfer rates, and survey scores, with Command dashboards for the remaining operational metrics.",
        ),
        (
            "13 call categories identified in Bid Line 4",
            "12 call categories identified in Bid Line 4",
        ),
    ]

    n = replace_all(doc, pairs)
    print(f"Word paragraph/cell replacements (pass 1): {n}")

    # Cover table identity
    cover = None
    for t in doc.tables:
        if t.rows and "RFP Number" in t.rows[0].cells[0].text:
            cover = t
            break
    if cover:
        for row in cover.rows:
            label = row.cells[0].text.strip()
            if label == "Respondent":
                set_cell_text(row.cells[1], ENTITY)
            elif label == "Prepared For":
                set_cell_text(
                    row.cells[1],
                    "Sharene Marquez, Purchasing Agent  Sharene.Marquez@kcpd.org | (816) 234-5340",
                )
            elif label == "Proposed Solution":
                set_cell_text(
                    row.cells[1],
                    "Rapid Cortex Command + Call Assist — AI-Assisted Non-Emergency Call Management",
                )

    # Qualifications table
    qual = None
    for t in doc.tables:
        if t.rows and "Qualification Requirement" in t.rows[0].cells[0].text:
            qual = t
            break
    if qual:
        updates = {
            "Minimum five years in public safety technology": (
                f"Addressed — {ENTITY} (Columbus, Georgia) is a public safety technology company. "
                f"The Rapid Cortex platform and Call Assist module are purpose-built for PSAP operations. "
                f"Live scenario demonstration is offered as part of evaluation."
            ),
            "Existing U.S. PSAP customers": (
                "Addressed — Rapid Cortex will provide law enforcement / PSAP references via the "
                "Reference Information Sheet. Evaluation includes a live Call Assist demonstration "
                "against KCPD-style scenarios."
            ),
            "Experience with agencies serving populations over 500,000": (
                "Addressed — Command Large T3 is the plan tier for 51–65 telecommunicator seats and "
                "large-metro non-emergency volume. Architecture (multi-AZ AWS, tenant isolation, "
                "Mission-Critical support) is sized for KCPD's operational environment."
            ),
            "References from at least three law enforcement agencies": (
                "Met — Three law enforcement agency references provided in Appendix A / the Reference "
                "Information Sheet, available for direct contact by the KCPD evaluation team."
            ),
            "Proven AI deployments in public safety": (
                "Addressed — AI-assisted intake, triage, summarization, and fail-closed CAD handoff "
                "are implemented in the Rapid Cortex / Call Assist platform. Live demonstration uses "
                "the seeded scenario runner against the Safety / Triage / Intake pipeline."
            ),
            "CAD platform agnostic": (
                f"Met — CAD abstraction layer with adapters for {CAD_ADAPTERS}. "
                f"PremierOne is the KCPD production path; other adapters share the same fail-closed gates."
            ),
        }
        for row in qual.rows[1:]:
            req = row.cells[0].text.strip()
            for key, val in updates.items():
                if key in req:
                    set_cell_text(row.cells[1], val)
                    break

    # Pricing table (Line / Component / Type / Price / Notes)
    price_tbl = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Line" and "Price" in t.rows[0].cells[3].text:
            price_tbl = t
            break
    if price_tbl:
        line_notes = {
            "1": (
                money(P[1]),
                f"Includes Command Large T3 platform base ({money(CMD_BASE)}/yr) + Call Assist module ({money(CALL_ASSIST)}/yr). 35% proposal discount.",
            ),
            "3": (money(P[3]), "AI Triage Workflows — 25+ classification rules (Premium)"),
            "5": (
                money(P[5]),
                "NLP/LLM engine, confidence scoring, model redundancy, summarization, Comprehend sentiment / Contact Lens distress",
            ),
            "6": (
                money(P[6]),
                "Enhanced transcription, EN+ES, live interpreter bridge, TTY/TDD via Connect + SMS fallback",
            ),
            "7": (
                money(P[7]),
                "Voice (included), SMS link generation, caller video upload via existing RC media",
            ),
            "8": (
                money(P[8]),
                "Callback number capture, after-hours offer, SMS status notifications; outbound campaigns configured in implementation",
            ),
            "11": (money(P[11]), CAD_NOTE),
            "14": (money(P[14]), RMS_NOTE),
            "16": (
                money(P[16]),
                "CJIS-aligned controls review and documentation package. MFA, encryption, RBAC, audit logging included in platform. Not a CJIS/SOC 2/FedRAMP certificate.",
            ),
            "18": (money(P[18]), "Extended retention storage (7-yr audio, transcript); Missouri Sunshine Law export workflows; chain of custody controls"),
            "19": (
                money(P[19]),
                "Command includes supervisor QA tools. This line: AI QA scoring, playback, keyword search, false-transfer tracking, escalation statistics, performance dashboards",
            ),
            "23": (money(P[23]), "Onsite dispatcher + supervisor training; remote admin training; CJIS training documentation; train-the-trainer; custom materials"),
            "24": (money(P[24]), "Project manager; implementation schedule; testing; UAT; parallel operations; cutover; rollback; go-live support; post-implementation review"),
            "25": (money(P[25]), "Mission-Critical support: 24×7 technical support, named account manager, updates, patches, QBRs, annual health checks"),
        }
        for row in price_tbl.rows[1:]:
            line = row.cells[0].text.strip()
            if line in line_notes:
                price, note = line_notes[line]
                set_cell_text(row.cells[3], price)
                set_cell_text(row.cells[4], note)

    # Contract year summary
    yr = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Period":
            yr = t
            break
    if yr:
        mapping = {
            "One-Time Implementation & Integration": money(ONE_TIME),
            "Annual Subscription (Year 1)": money(ANNUAL),
            "Total Year 1 Investment": money(YEAR1),
            "Year 2+ Annual Recurring": money(YEAR2),
        }
        for row in yr.rows[1:]:
            period = row.cells[0].text.strip()
            for key, val in mapping.items():
                if period.startswith(key.split("(")[0].strip()) or key in period:
                    set_cell_text(row.cells[1], val)
                    break

    # Bid Line 4 / 12 vs 13 leftover
    extra = [
        (
            "Rapid Cortex automatically classifies each incoming call across all 12 categories specified in the RFP, plus additional configurable types:",
            "Rapid Cortex automatically classifies each incoming call across all 12 categories specified in the RFP, plus additional configurable types:",
        ),
        (
            "Assisted write-back: AI-populated incident drafts (type, location, notes, priority) are presented to the call taker for one-click confirmation before writing to CAD. Dispatchers retain full control of every CAD entry.",
            "Assisted write-back: AI-populated incident drafts (type, location, notes, priority) are presented to the call taker for confirmation before writing to CAD. Dispatchers retain full control of every CAD entry. Production write-back is fail-closed until KCPD UAT; Rapid Cortex does not auto-dispatch.",
        ),
        (
            "Rapid Cortex connects to KCPD's non-emergency telephony via Amazon Connect, Rapid Cortex's preferred telephony backbone for PSAP deployments. The integration supports SIP trunk connection to KCPD's existing telephony infrastructure, with ANI/ALI passthrough preserved through the connection. No telephony hardware replacement is required.",
            "Rapid Cortex connects to KCPD's non-emergency telephony via Amazon Connect. The integration supports SIP trunk connection to KCPD's existing telephony infrastructure, with ANI/ALI passthrough preserved. Call Assist does not replace KCPD's 911 CPE or phone system. No telephony hardware replacement is required.",
        ),
        (
            "Rapid Cortex's architecture is designed and operated in alignment with the FBI CJIS Security Policy. Compliance elements include:",
            "Rapid Cortex is designed with CJIS Security Policy–aligned controls. Control elements include:",
        ),
    ]
    n2 = replace_all(doc, extra)

    doc.save(str(DOCX))
    print(f"Word extra replacements: {n2}")
    print(f"Saved {DOCX}")


def verify() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ps = wb["Pricing Schedule"]
    print("VERIFY Excel D7", ps["D7"].value, "D36", ps["D36"].value, "B3", ps["B3"].value)
    doc = Document(str(DOCX))
    blob = "\n".join(p.text for p in doc.paragraphs)
    for needle in ("Sharen.Marquez", "$2,503,000", "$648,000", "CommandAI", "SOC 2 Type II alignment"):
        print(f"  leftover {needle!r}:", needle in blob)
    print("  has Apps on Demand:", ENTITY in blob or any(ENTITY in t.rows[0].cells[1].text for t in doc.tables if t.rows))
    # pricing table
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Line":
            print("  Line 1 price cell:", t.rows[1].cells[3].text)
            print("  Line 11 price cell:", t.rows[11].cells[3].text)
            break


if __name__ == "__main__":
    print("ANNUAL", ANNUAL, "ONE_TIME", ONE_TIME, "YEAR1", YEAR1, "YEAR2", YEAR2)
    print("CAD_MAINT", CAD_MAINT, "RMS_MAINT", RMS_MAINT)
    print("P11 components check", 9750 + 9750 + 11375 + 16250 + 81250 + 13000 + 6500 + 9750 + 23563)
    update_xlsx()
    update_docx()
    verify()
